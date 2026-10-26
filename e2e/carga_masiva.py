"""
Carga masiva desde Excel (INT-01) en un navegador real: se descarga la plantilla desde la app, se llena con
openpyxl (como lo haría alguien en Excel), se sube con errores, se corrige y se carga.
Requiere los datos de demostración recién cargados y el frontend compilado (npm run build && npm run preview).
Uso: python e2e/carga_masiva.py <contraseña-admin>
"""
import sys, re, os, tempfile, datetime
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright, expect

BASE = os.environ.get('CMMS_URL', 'http://localhost:4173')
ADMIN_PASS = sys.argv[1]
SP = os.environ.get('SHOTS', tempfile.mkdtemp(prefix='cmms-carga-'))
errors, results = [], []

def ok(n):
    results.append(n); print('OK  ', n)

def track(page):
    page.on('pageerror', lambda e: errors.append(f'{page.url}: {e}'))
    page.on('console', lambda m: m.type == 'error' and 'Failed to load resource' not in m.text and errors.append(f'{page.url}: {m.text}'))

def login(page, u, c):
    page.goto(BASE + '/ingresar')
    page.get_by_label('Usuario').fill(u)
    page.get_by_label('Contraseña o PIN').fill(c)
    page.get_by_role('button', name='Ingresar').click()
    expect(page).to_have_url(re.compile('/ordenes'))

def descargar_plantilla(page, nombre):
    with page.expect_download() as d:
        page.get_by_role('button', name=re.compile('Plantilla de')).click()
    dl = d.value
    assert dl.suggested_filename == nombre, dl.suggested_filename
    ruta = os.path.join(SP, nombre)
    dl.save_as(ruta)
    return ruta

def llenar(ruta, filas, salida):
    wb = load_workbook(ruta)
    ws = wb.worksheets[0]
    for i, f in enumerate(filas, start=2):
        for j, v in enumerate(f, start=1):
            ws.cell(row=i, column=j, value=v)
    wb.save(salida)
    return salida

with sync_playwright() as p:
    b = p.chromium.launch()

    # 1) Admin en el PC: máquinas
    pc = b.new_context(viewport={'width': 1366, 'height': 900}, locale='es-CO', timezone_id='America/Bogota', accept_downloads=True)
    page = pc.new_page(); track(page)
    login(page, 'admin', ADMIN_PASS)
    page.goto(BASE + '/maquinas')
    page.get_by_role('link', name='Desde Excel').click()
    expect(page).to_have_url(re.compile(r'/configuracion/carga-masiva\?tipo=maquinas'))
    expect(page.get_by_role('button', name='Máquinas', exact=True)).to_have_attribute('aria-pressed', 'true')
    plantilla = descargar_plantilla(page, 'plantilla-maquinas.xlsx')
    wb = load_workbook(plantilla)
    assert [c.value for c in wb.worksheets[0][1]][:3] == ['Código *', 'Nombre *', 'Área *'], [c.value for c in wb.worksheets[0][1]]
    assert 'Instrucciones' in wb.sheetnames
    ok('admin descarga la plantilla de máquinas desde la lista de máquinas (con hoja de instrucciones)')

    # Encabezados: Código, Nombre, Área, Ubicación, Fabricante, Modelo, Serie, Código SAP, Fecha de compra, Fin de garantía, Estado, Notas
    filas_mal = [
        ['TAJ-001', 'Tajadora de jamón', 'Proceso', 'Sala 2', 'Bizerba', 'A550', 'S-889', 'SAP-1001', datetime.date(2023, 3, 15), datetime.date(2026, 3, 15), 'Operativo', ''],
        ['MOL-001', 'Molino repetido', 'Proceso'],
        ['EMP-001', 'Empacadora', '', 'Empaque'],
        ['TAJ-001', 'Tajadora otra vez', 'Proceso'],
        ['CAL-001', 'Caldera', 'Servicios', 'Cuarto de calderas', '', '', '', '', '30/02/2020'],
    ]
    page.get_by_label('Archivo de Excel').set_input_files(llenar(plantilla, filas_mal, os.path.join(SP, 'maquinas-con-errores.xlsx')))
    expect(page.locator('.import-errors tbody tr')).to_have_count(4)
    expect(page.locator('.import-errors')).to_contain_text('MOL-001 ya existe en el sistema')
    expect(page.locator('.import-errors')).to_contain_text('TAJ-001 está repetido (también en la fila 2)')
    expect(page.locator('.import-errors tbody tr', has_text='Área')).to_contain_text('4')
    expect(page.locator('.import-errors')).to_contain_text('La fecha no existe')
    expect(page.get_by_role('button', name=re.compile(r'^Cargar \d'))).to_have_count(0)
    page.screenshot(path=f'{SP}/int-01-errores.png', full_page=True)
    ok('archivo con 4 errores: la app muestra fila, columna y problema y no deja cargar')

    filas_bien = [
        filas_mal[0],
        ['EMP-001', 'Empacadora al vacío', 'Empaque', 'Sala de empaque', 'Multivac', 'R245', '', '', '10/01/2021', '', '', 'Revisar sellos'],
        ['CAL-001', 'Caldera', 'Servicios', 'Cuarto de calderas', '', '', '', '', '28/02/2020'],
        ['BAS-001', 'Báscula antigua', 'Despachos', '', '', '', '', '', '', '', 'Dado de baja'],
    ]
    page.get_by_label('Archivo de Excel').set_input_files(llenar(plantilla, filas_bien, os.path.join(SP, 'maquinas-corregido.xlsx')))
    expect(page.locator('.import-errors')).to_have_count(0)
    expect(page.locator('.kpi', has_text='Listas para cargar')).to_contain_text('4')
    expect(page.locator('.kpi', has_text='Listas para cargar')).to_contain_text('1 dada de baja')
    expect(page.locator('table tbody tr')).to_have_count(4)
    expect(page.locator('table tbody tr').first).to_contain_text('15/03/2023')
    page.screenshot(path=f'{SP}/int-02-vista-previa.png', full_page=True)
    page.get_by_role('button', name='Cargar 4 máquinas').click()
    expect(page.locator('.alert-ok')).to_contain_text('Se cargaron 4 máquinas')
    page.get_by_role('link', name='Ver máquinas').click()
    for c in ['TAJ-001', 'EMP-001', 'CAL-001', 'BAS-001']:
        expect(page.get_by_text(c, exact=True)).to_be_visible()
    ok('archivo corregido: vista previa, carga de 4 máquinas y aparecen en la lista')

    # 2) Admin en el celular: repuestos con stock inicial
    cel = b.new_context(**p.devices['Pixel 7'], locale='es-CO', timezone_id='America/Bogota', accept_downloads=True)
    m = cel.new_page(); track(m)
    login(m, 'admin', ADMIN_PASS)
    m.goto(BASE + '/configuracion')
    m.get_by_role('link', name=re.compile('Carga masiva desde Excel')).click()
    m.get_by_role('button', name='Repuestos', exact=True).click()
    plantilla_r = descargar_plantilla(m, 'plantilla-repuestos.xlsx')
    # Código, Descripción, Unidad, Ubicación, Código SAP, Stock mínimo, Stock inicial, Costo unitario
    filas_r = [
        ['RP-COR-B42', 'Correa en V B42', 'und', 'Estante C-1', '', 2, 6, 23500],
        ['RP-ACE-68', 'Aceite hidráulico ISO 68', 'L', 'Bodega lubricantes', 'SAP-5521', 10, 40.5, '9.800'],
        ['RP-SEN-PT100', 'Sensor PT100', 'UND', 'Estante D-4', '', 1, '', ''],
    ]
    m.get_by_label('Archivo de Excel').set_input_files(llenar(plantilla_r, filas_r, os.path.join(SP, 'repuestos.xlsx')))
    resumen = m.locator('.kpi', has_text='Listas para cargar')
    expect(resumen).to_contain_text('3')
    expect(resumen).to_contain_text('2 con stock inicial')
    expect(resumen).to_contain_text('46,5 unidades')
    expect(resumen).to_contain_text('537.900')  # 6 × 23.500 + 40,5 × 9.800
    assert m.evaluate('document.documentElement.scrollWidth <= window.innerWidth + 1'), 'desplazamiento horizontal'
    m.screenshot(path=f'{SP}/int-03-repuestos-celular.png', full_page=True)
    m.get_by_role('button', name='Cargar 3 repuestos').click()
    expect(m.locator('.alert-ok')).to_contain_text('Se cargaron 3 repuestos')
    ok('celular: repuestos con stock inicial; totales (46,5 unidades, $ 537.900) antes de cargar')

    m.get_by_role('link', name='Ver inventario').click()
    m.get_by_role('link', name=re.compile('RP-ACE-68')).click()
    expect(m.get_by_text('Saldo inicial', exact=True)).to_be_visible()
    expect(m.locator('body')).to_contain_text('40,5')
    ok('el aceite queda con 40,5 L y su "Saldo inicial" en el kárdex')

    # 3) Bitácora y el mismo archivo otra vez
    page.goto(BASE + '/configuracion')
    expect(page.locator('.audit').first).to_contain_text('Carga masiva desde Excel')
    expect(page.locator('.audit').first).to_contain_text('repuestos.xlsx')
    page.goto(BASE + '/configuracion/carga-masiva?tipo=repuestos')
    page.get_by_label('Archivo de Excel').set_input_files(os.path.join(SP, 'repuestos.xlsx'))
    expect(page.locator('.import-errors tbody tr')).to_have_count(3)
    expect(page.locator('.import-errors')).to_contain_text('RP-COR-B42 ya existe en el sistema')
    ok('queda en la bitácora; subir el mismo archivo otra vez no duplica (3 códigos ya existen)')

    # 4) Archivo que no es Excel
    falso = os.path.join(SP, 'no-es-excel.xlsx')
    open(falso, 'wb').write(b'%PDF-1.4 esto no es excel')
    page.get_by_label('Archivo de Excel').set_input_files(falso)
    expect(page.locator('.alert-bad, .error-box').first).to_contain_text('Solo se aceptan archivos de Excel (.xlsx)')
    ok('un archivo que no es Excel se rechaza con un mensaje claro')

    # 5) El técnico no tiene la opción
    t = b.new_context(**p.devices['iPhone 13'], locale='es-CO', timezone_id='America/Bogota').new_page(); track(t)
    login(t, 'tecnico', '4826')
    t.goto(BASE + '/maquinas')
    expect(t.get_by_role('link', name='Desde Excel')).to_have_count(0)
    t.goto(BASE + '/configuracion/carga-masiva')
    expect(t.get_by_role('button', name=re.compile('Plantilla de'))).to_have_count(0)
    ok('técnico: sin acceso a la carga masiva')
    b.close()

print('\n', len(results), 'verificaciones OK. Capturas en', SP)
print('Errores de consola:', errors or 'ninguno')
